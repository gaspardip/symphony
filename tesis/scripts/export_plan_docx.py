from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "tesis" / "propuesta" / "plan.md"
OUTPUT = ROOT / "tesis" / "entrega" / "plan-tesis-proyecto-final-integrador.docx"

TITLE = "Arquitectura de runtime con proof gates para delivery autónomo de software basado en issues"
SUBTITLE = "Plan de Tesis - Proyecto Final Integrador"
FOOTER_LEFT = "Dip"

FONT = "Times New Roman"
MONO = "Courier New"
NAVY = RGBColor(23, 54, 93)
BLUE = RGBColor(31, 78, 121)
INK = RGBColor(25, 25, 25)
MUTED = RGBColor(85, 85, 85)
WHITE = RGBColor(255, 255, 255)
LIGHT_GRAY = "F4F6F8"
TABLE_STRIPE = "F8FAFC"
CONTENT_WIDTH_DXA = 9072


def set_run_font(
    run,
    *,
    size: float = 11,
    name: str = FONT,
    color: RGBColor = INK,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_paragraph_format(
    paragraph,
    *,
    line_spacing: float = 1.5,
    before: float = 0,
    after: float = 4,
    keep_with_next: bool = False,
) -> None:
    paragraph.paragraph_format.line_spacing = line_spacing
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.keep_with_next = keep_with_next


def paragraph_border(paragraph, *, bottom: str | None = None, top: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for side, color in [("bottom", bottom), ("top", top)]:
        if not color:
            continue
        elem = OxmlElement(f"w:{side}")
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "8")
        elem.set(qn("w:space"), "4")
        elem.set(qn("w:color"), color)
        p_bdr.append(elem)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, *, top: int = 90, start: int = 110, bottom: int = 90, end: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], *, indent_dxa: int = 120) -> None:
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(i, len(widths) - 1)])


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    result = OxmlElement("w:t")
    result.text = "1"
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)
    run._r.append(result)
    run._r.append(fld_char3)


def configure_section(section) -> None:
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.25)
    section.bottom_margin = Cm(2.25)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.35)
    section.header_distance = Cm(1.05)
    section.footer_distance = Cm(1.05)

    hp = section.header.paragraphs[0]
    hp.text = ""
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(hp, line_spacing=1.0, after=0)
    run = hp.add_run(TITLE)
    set_run_font(run, size=8.5, color=MUTED)

    fp = section.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_format(fp, line_spacing=1.0, after=0)
    fp.paragraph_format.tab_stops.add_tab_stop(Cm(16.2), WD_TAB_ALIGNMENT.RIGHT)
    paragraph_border(fp, top="BFBFBF")
    left = fp.add_run(FOOTER_LEFT)
    set_run_font(left, size=8.5, color=MUTED)
    tab = fp.add_run("\t")
    set_run_font(tab, size=8.5, color=MUTED)
    label = fp.add_run("Página ")
    set_run_font(label, size=8.5, color=MUTED)
    add_page_number(fp)
    for r in fp.runs:
        set_run_font(r, size=8.5, color=MUTED)


def configure_document(doc: Document) -> None:
    configure_section(doc.sections[0])
    styles = doc.styles
    for name in ["Normal", "Body Text"]:
        if name in styles:
            style = styles[name]
            style.font.name = FONT
            style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
            style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
            style.font.size = Pt(11)
            style.font.color.rgb = INK
            style.paragraph_format.line_spacing = 1.5
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(4)
    heading_tokens = {
        "Heading 1": (16, NAVY, 12, 6),
        "Heading 2": (14, BLUE, 10, 4),
        "Heading 3": (12, NAVY, 8, 3),
        "Heading 4": (11, BLUE, 6, 2),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        if name in styles:
            style = styles[name]
            style.font.name = FONT
            style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
            style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
            style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = color
            style.paragraph_format.line_spacing = 1.2
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True


def add_inline(
    paragraph,
    text: str,
    *,
    default_size: float = 11,
    default_color: RGBColor = INK,
    bold_all: bool = False,
    italic_all: bool = False,
) -> None:
    text = text.replace("\\|", "|")
    buf: list[str] = []
    bold = bold_all
    italic = italic_all
    i = 0

    def flush() -> None:
        nonlocal buf
        if not buf:
            return
        run = paragraph.add_run("".join(buf))
        set_run_font(run, size=default_size, color=default_color, bold=bold, italic=italic)
        buf = []

    while i < len(text):
        if text.startswith("`", i):
            end = text.find("`", i + 1)
            if end != -1:
                flush()
                run = paragraph.add_run(text[i + 1 : end])
                set_run_font(run, size=default_size - 0.5, name=MONO, color=default_color, bold=bold, italic=italic)
                i = end + 1
                continue
        if text.startswith("***", i):
            flush()
            bold = not bold
            italic = not italic
            i += 3
            continue
        if text.startswith("**", i):
            flush()
            bold = not bold
            i += 2
            continue
        if text.startswith("*", i):
            flush()
            italic = not italic
            i += 1
            continue
        buf.append(text[i])
        i += 1
    flush()


def add_para(
    doc: Document,
    text: str = "",
    *,
    style: str | None = None,
    alignment=None,
    size: float = 11,
    color: RGBColor = INK,
    bold: bool = False,
    italic: bool = False,
    before: float = 0,
    after: float = 4,
    line_spacing: float = 1.5,
    keep_with_next: bool = False,
) -> None:
    p = doc.add_paragraph(style=style)
    if alignment is not None:
        p.alignment = alignment
    set_paragraph_format(p, line_spacing=line_spacing, before=before, after=after, keep_with_next=keep_with_next)
    add_inline(p, text, default_size=size, default_color=color, bold_all=bold, italic_all=italic)


def add_spacer(doc: Document, points: float) -> None:
    p = doc.add_paragraph()
    set_paragraph_format(p, line_spacing=1.0, before=0, after=points)


def add_cover_band(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_geometry(table, [CONTENT_WIDTH_DXA], indent_dxa=0)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "17365D")
    set_cell_margins(cell, top=180, bottom=180, start=180, end=180)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(p, line_spacing=1.1, after=0)
    add_inline(p, text, default_size=10.5, default_color=WHITE, bold_all=True)
    add_spacer(doc, 12)


def add_metadata_table(doc: Document) -> None:
    rows = [
        ("Autor", "Gaspar Dip"),
        ("Matrícula", "49012185"),
        ("Correo", "dip.gaspar@gmail.com"),
        ("Carrera", "Ingeniería en Informática"),
        ("Asignatura", "Tesis de Grado - Proyecto Final Integrador"),
        ("Director de tesis", "Dr. Jorge Salvador Ierache"),
        ("Codirector", "No corresponde"),
        ("Fecha", "21/05/2026"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_geometry(table, [2300, CONTENT_WIDTH_DXA - 2300], indent_dxa=0)
    for i, (label, value) in enumerate(rows):
        label_cell = table.cell(i, 0)
        value_cell = table.cell(i, 1)
        set_cell_shading(label_cell, LIGHT_GRAY)
        set_cell_margins(label_cell, top=100, bottom=100, start=150, end=120)
        set_cell_margins(value_cell, top=100, bottom=100, start=150, end=120)
        set_cell_text(label_cell, label, bold=True, size=9.5, color=NAVY)
        set_cell_text(value_cell, value, size=9.5)
    add_spacer(doc, 14)


def add_cover(doc: Document) -> None:
    add_spacer(doc, 34)
    add_para(
        doc,
        "Escuela Superior de Ingeniería, Informática y Cs. Agroalimentarias",
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        size=11,
        color=MUTED,
        bold=True,
        after=8,
        line_spacing=1.1,
    )
    add_para(
        doc,
        SUBTITLE,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        size=17,
        color=NAVY,
        bold=True,
        after=14,
        line_spacing=1.1,
    )
    add_cover_band(doc, TITLE)
    add_para(
        doc,
        "Tema: ejecución controlada, observable y verificable de agentes de software sobre repositorios.",
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        size=10.5,
        color=MUTED,
        after=18,
        line_spacing=1.2,
    )
    add_metadata_table(doc)
    add_para(
        doc,
        "Documento preparado para la entrega de Blackboard del Plan de Tesis - Proyecto Final Integrador.",
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        size=9.5,
        color=MUTED,
        italic=True,
        after=0,
        line_spacing=1.2,
    )


def set_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    size: float = 8.2,
    color: RGBColor = INK,
    align=None,
) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    set_paragraph_format(p, line_spacing=1.08, after=0)
    add_inline(p, text, default_size=size, default_color=color, bold_all=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def parse_table(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for raw in lines:
        line = raw.strip()
        if line.startswith("|"):
            line = line[1:]
        if line.endswith("|"):
            line = line[:-1]
        cells = [c.strip() for c in line.split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c.strip()) for c in cells):
            continue
        rows.append(cells)
    return rows


def column_widths(column_count: int) -> list[int]:
    patterns = {
        1: [CONTENT_WIDTH_DXA],
        2: [2600, CONTENT_WIDTH_DXA - 2600],
        3: [2300, 1500, CONTENT_WIDTH_DXA - 3800],
        4: [1700, 2250, 2250, CONTENT_WIDTH_DXA - 6200],
        5: [1450, 2050, 1400, 2050, CONTENT_WIDTH_DXA - 6950],
    }
    if column_count in patterns:
        return patterns[column_count]
    base = CONTENT_WIDTH_DXA // column_count
    widths = [base] * column_count
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def add_table(doc: Document, table_lines: list[str]) -> None:
    rows = parse_table(table_lines)
    if not rows:
        return
    max_cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=max_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    widths = column_widths(max_cols)
    set_table_geometry(table, widths)
    for i, row in enumerate(rows):
        if i == 0:
            set_repeat_table_header(table.rows[i])
        for j in range(max_cols):
            cell = table.cell(i, j)
            if i == 0:
                set_cell_shading(cell, "17365D")
                set_cell_text(cell, row[j] if j < len(row) else "", bold=True, size=8.1, color=WHITE)
            else:
                if i % 2 == 0:
                    set_cell_shading(cell, TABLE_STRIPE)
                set_cell_text(cell, row[j] if j < len(row) else "", size=7.8)
            set_cell_margins(cell, top=90, bottom=90, start=100, end=100)
    add_spacer(doc, 8)


def add_caption(doc: Document, text: str) -> None:
    clean = text.strip()
    if clean.startswith("**") and clean.endswith("**"):
        clean = clean[2:-2]
    p = doc.add_paragraph()
    set_paragraph_format(p, line_spacing=1.15, before=2, after=5, keep_with_next=True)
    paragraph_border(p, bottom="D9E2F3")
    add_inline(p, clean, default_size=9.5, default_color=NAVY, bold_all=True)


def add_markdown_body(doc: Document, text: str, *, start_on_new_page: bool = False) -> None:
    lines = text.splitlines()
    i = 0
    in_code = False
    code_lines: list[str] = []
    first_body_paragraph = True

    def mark_body_start(paragraph) -> None:
        nonlocal first_body_paragraph
        if first_body_paragraph and start_on_new_page:
            paragraph.paragraph_format.page_break_before = True
        first_body_paragraph = False

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                set_paragraph_format(p, line_spacing=1.0, before=3, after=6)
                run = p.add_run("\n".join(code_lines))
                set_run_font(run, size=8.5, name=MONO, color=INK)
                code_lines = []
                in_code = False
            else:
                in_code = True
                code_lines = []
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not stripped:
            i += 1
            continue
        if stripped == "---":
            p = doc.add_paragraph()
            mark_body_start(p)
            set_paragraph_format(p, line_spacing=1.0, before=4, after=7)
            paragraph_border(p, bottom="BFBFBF")
            i += 1
            continue
        if stripped.startswith("|") and "|" in stripped[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_table(doc, table_lines)
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            level = min(len(heading.group(1)), 4)
            content = heading.group(2).strip()
            p = doc.add_paragraph(style=f"Heading {level}")
            mark_body_start(p)
            set_paragraph_format(
                p,
                line_spacing=1.15,
                before=10 if level <= 2 else 6,
                after=5 if level <= 2 else 3,
                keep_with_next=True,
            )
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(
                p,
                content,
                default_size=16 if level == 1 else 14 if level == 2 else 12,
                default_color=NAVY if level <= 2 else BLUE,
                bold_all=True,
            )
            i += 1
            continue

        if stripped.startswith("**Tabla ") or stripped.startswith("**Figura "):
            add_caption(doc, stripped)
            i += 1
            continue

        unordered = re.match(r"^[-*]\s+(.*)$", stripped)
        if unordered:
            p = doc.add_paragraph(style="List Bullet")
            mark_body_start(p)
            set_paragraph_format(p, line_spacing=1.3, before=0, after=3)
            add_inline(p, unordered.group(1), default_size=10.8)
            i += 1
            continue

        ordered = re.match(r"^\d+\.\s+(.*)$", stripped)
        if ordered:
            p = doc.add_paragraph(style="List Number")
            mark_body_start(p)
            set_paragraph_format(p, line_spacing=1.3, before=0, after=3)
            add_inline(p, ordered.group(1), default_size=10.8)
            i += 1
            continue

        paragraph_parts = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt:
                i += 1
                break
            if (
                nxt == "---"
                or nxt.startswith("#")
                or nxt.startswith("|")
                or nxt.startswith("```")
                or nxt.startswith("**Tabla ")
                or nxt.startswith("**Figura ")
                or re.match(r"^[-*]\s+", nxt)
                or re.match(r"^\d+\.\s+", nxt)
            ):
                break
            paragraph_parts.append(nxt)
            i += 1
        p = doc.add_paragraph()
        mark_body_start(p)
        set_paragraph_format(p, line_spacing=1.35, after=5)
        add_inline(p, " ".join(paragraph_parts), default_size=11, default_color=INK)


def main() -> None:
    md = SOURCE.read_text(encoding="utf-8")
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_markdown_body(doc, md, start_on_new_page=True)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
